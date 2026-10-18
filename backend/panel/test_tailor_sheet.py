"""Tailor sheet (Prompt 25): built as a background job, one sheet per school, fixed grade order."""

import io
import random
import shutil
import tempfile
import uuid
from decimal import Decimal

from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import User
from accounts.serializers import build_tokens_for_user
from catalog.models import Category, Product, ProductVariant
from orders.models import Order, OrderItem
from schools.models import City, Grade, School, Student

from .models import ExportJob

URL = "/api/panel/tailor-sheets/"


def login(user):
    client = APIClient()
    client.credentials(HTTP_AUTHORIZATION=f"Bearer {build_tokens_for_user(user)['access']}")
    return client


class TailorSheetTests(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._media = tempfile.mkdtemp(prefix="tailor-sheet-")
        cls._override = override_settings(MEDIA_ROOT=cls._media)
        cls._override.enable()

    @classmethod
    def tearDownClass(cls):
        cls._override.disable()
        shutil.rmtree(cls._media, ignore_errors=True)
        super().tearDownClass()

    @classmethod
    def setUpTestData(cls):
        rng = random.Random(25)
        cls.surat = City.objects.create(name="Surat T", code="SURT")
        cls.far = City.objects.create(name="Far T", code="FART")
        # Names chosen so A-Z order differs from creation order.
        cls.schools = [
            School.objects.create(city=cls.surat, name=name, code=f"T{i}")
            for i, name in enumerate(["Vidya Mandir", "Alpha Academy", "Kendriya", "St Mary's [Main]", "Bright Future"])
        ]
        cls.far_school = School.objects.create(city=cls.far, name="Far School", code="TFAR")
        grades = list(Grade.objects.order_by("sort_order"))
        if not grades:
            names = ["Nursery", "Junior KG", "Senior KG"] + [f"Class {n}" for n in range(1, 13)]
            grades = [Grade.objects.create(name=n, sort_order=i) for i, n in enumerate(names, start=1)]
        cls.grade_order = [g.name for g in grades]

        uniform = Category.objects.create(name="Uniform T", slug="uniform")
        shoes = Category.objects.create(name="School Shoes T", slug="school-shoes")
        accessories = Category.objects.create(name="Uniform Accessories T", slug="uniform-accessories")

        def product(name, category, sizes, prefix, kind=None):
            item = Product.objects.create(
                category=category, name=name, cost_price=1, selling_price=2, product_type=kind
            )
            return [ProductVariant.objects.create(product=item, size=s, sku=f"{prefix}-{s}") for s in sizes]

        shirt = product("Half Shirt", uniform, ["26", "28", "30", "32"], "T-SH")
        trouser = product("Grey Trousers", uniform, ["22", "24", "26"], "T-TR")
        shoe = product("Black Shoes", shoes, ["UK-2", "UK-3", "UK-4"], "T-SHOE")
        tie = product("Striped Tie", accessories, ["Standard"], "T-TIE", "TIE")

        cls.boss = User.objects.create_user(username="t_boss", password="x", role="BOSS")
        cls.admin = User.objects.create_user(username="t_admin", password="x", role="ADMIN", city=cls.surat)
        cls.school_admin = User.objects.create_user(
            username="t_sa", password="x", role="SCHOOL_ADMIN", school=cls.schools[0], city=cls.surat
        )
        parent = User.objects.create_user(username="t_parent", password="x", role="PARENT")

        students = []
        for school in cls.schools + [cls.far_school]:
            for n in range(30):
                grade = rng.choice(grades)
                students.append(
                    Student(
                        name=f"Child {n:02d} {school.code}",
                        gr_number=f"{school.code}-{n}",
                        grade=grade,
                        class_name=grade.name,
                        section=rng.choice("AB"),
                        gender=rng.choice(["MALE", "FEMALE"]),
                        school=school,
                        city=school.city,
                    )
                )
        Student.objects.bulk_create(students)
        students = list(Student.objects.select_related("grade"))

        orders, items = [], []
        cls.expected = {}  # order number -> expected uniform rows (product, size, height, shoulder, edited)
        for n in range(2000):
            student = students[n % len(students)]
            order = Order(
                order_number=f"ORD-T{n:05d}",
                idempotency_key=str(uuid.uuid4()),
                placed_by=parent,
                placed_by_role="PARENT",
                student=student,
                student_class=student.class_name,
                school_id=student.school_id,
                city_id=student.city_id,
                status=Order.Status.PLACED,
                payment_method=Order.PaymentMethod.CASH_ON_DELIVERY,
                subtotal=Decimal("8"),
                total=Decimal("8"),
            )
            orders.append(order)
            height = Decimal(rng.randint(100, 170))
            shoulder = Decimal(str(round(rng.uniform(24, 40), 1)))
            edited = n % 7 == 0
            outside = n % 11 == 0
            chosen_shirt, chosen_trouser, chosen_shoe = rng.choice(shirt), rng.choice(trouser), rng.choice(shoe)
            for variant, measured in ((chosen_shirt, True), (chosen_trouser, True)):
                items.append(
                    OrderItem(
                        order=order, variant=variant, category=uniform, quantity=1,
                        unit_price_snapshot=2, unit_cost_snapshot=1, chosen_size=variant.size,
                        fit_height_cm=height, fit_weight_kg=Decimal("30.0"),
                        fit_shoulder_cm=shoulder, fit_source="CAMERA", fit_shoulder_edited=edited,
                        recommendation_status="ABOVE_CHART" if outside else "FIT",
                    )
                )
            items.append(OrderItem(order=order, variant=chosen_shoe, category=shoes, quantity=1,
                                   unit_price_snapshot=2, unit_cost_snapshot=1, chosen_size=chosen_shoe.size))
            items.append(OrderItem(order=order, variant=tie[0], category=accessories, quantity=1,
                                   unit_price_snapshot=2, unit_cost_snapshot=1, chosen_size="Standard"))
            cls.expected[order.order_number] = {
                "school": student.school_id,
                "rows": sorted([
                    ("Half Shirt", chosen_shirt.size), ("Grey Trousers", chosen_trouser.size),
                ]),
                "height": float(height),
                "shoulder": float(shoulder),
                "edited": "yes" if edited else "no",
                "shoe": chosen_shoe.size,
                "highlight": edited or outside,
            }
        Order.objects.bulk_create(orders, batch_size=500)
        saved = {o.order_number: o for o in Order.objects.filter(order_number__startswith="ORD-T")}
        for item in items:
            item.order = saved[item.order.order_number]
        OrderItem.objects.bulk_create(items, batch_size=1000)
        # A dispatched order is not on the default sheet.
        Order.objects.filter(order_number="ORD-T00001").update(status=Order.Status.DISPATCHED)

    def build(self, user, filters=None):
        res = login(user).post(URL, {"filters": filters or {}}, format="json")
        self.assertEqual(res.status_code, 202, res.data)
        job = ExportJob.objects.get(pk=res.data["id"])
        self.assertEqual(job.status, ExportJob.Status.READY, job.error_message)
        detail = login(user).get(f"{URL}{job.id}/")
        self.assertTrue(detail.data["download_url"])
        download = login(user).get(detail.data["download_url"])
        self.assertEqual(download.status_code, 200)
        return job, load_workbook(io.BytesIO(b"".join(download.streaming_content)))

    def test_boss_sheet_has_every_school_grades_in_order_and_matching_sizes(self):
        job, book = self.build(self.boss)
        school_sheets = [name for name in book.sheetnames if name != "Export info"]
        expected_names = sorted(s.name for s in self.schools + [self.far_school])
        expected_names = [n.replace("[", " ").replace("]", " ").strip() for n in expected_names]
        self.assertEqual(school_sheets, expected_names)  # one sheet per school, A to Z

        seen_orders = set()
        flagged = 0
        for name in school_sheets:
            sheet = book[name]
            rows = list(sheet.iter_rows(values_only=True))
            self.assertEqual(rows[0][0], "Order no.")
            classes = []
            for row in rows[1:]:
                if row[1] and str(row[1]).endswith(("child", "children")):
                    classes.append(row[0].rsplit(" - Section", 1)[0])
                    continue
                order = self.expected[row[0]]
                seen_orders.add(row[0])
                self.assertIn((row[5], row[6]), order["rows"])
                self.assertEqual((row[7], row[9], row[11], row[12], row[13]),
                                 (order["height"], order["shoulder"], order["edited"], order["shoe"], "Standard"))
            positions = [self.grade_order.index(c) for c in classes]
            self.assertEqual(positions, sorted(positions), name)  # fixed grade order, never A to Z
            flagged += sum(
                1 for cells in sheet.iter_rows(min_row=2)
                if cells[0].fill is not None and cells[0].fill.fgColor.rgb in ("00FFF2CC", "FFFFF2CC")
            )

        self.assertNotIn("ORD-T00001", seen_orders)  # dispatched
        self.assertEqual(len(seen_orders), 1999)
        self.assertEqual(job.row_count, 1999 * 2)
        expected_flagged = sum(2 for number, o in self.expected.items() if o["highlight"] and number != "ORD-T00001")
        self.assertEqual(flagged, expected_flagged)
        self.assertEqual(job.result["rows_to_check"], expected_flagged)

    def test_admin_gets_only_their_city(self):
        _, book = self.build(self.admin)
        self.assertNotIn("Far School", book.sheetnames)
        self.assertEqual(len([n for n in book.sheetnames if n != "Export info"]), 5)
        res = login(self.admin).post(URL, {"filters": {"schools": [str(self.far_school.id)]}}, format="json")
        self.assertEqual(res.status_code, 400)

    def test_schools_cannot_get_it(self):
        self.assertEqual(login(self.school_admin).post(URL, {"filters": {}}, format="json").status_code, 403)
        self.assertEqual(login(self.school_admin).get(URL).status_code, 403)

    def test_filters_one_school_one_grade_and_one_order(self):
        school = self.schools[1]
        grade = Student.objects.filter(school=school).values_list("class_name", flat=True).first()
        _, book = self.build(self.admin, {"schools": [str(school.id)], "grade": grade})
        self.assertEqual([n for n in book.sheetnames if n != "Export info"], [school.name])
        headers = [r[0] for r in book[school.name].iter_rows(min_row=2, values_only=True)
                   if r[1] and str(r[1]).endswith(("child", "children"))]
        self.assertTrue(headers and all(h.startswith(grade) for h in headers))

        job, book = self.build(self.boss, {"order": str(Order.objects.get(order_number="ORD-T00001").id)})
        self.assertEqual(job.row_count, 2)  # one order, even though it is dispatched
        self.assertIn("ord-t00001", job.filename)
