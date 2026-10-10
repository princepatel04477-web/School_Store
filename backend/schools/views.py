import csv
import io

from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.models import User
from common.cache_mixin import VersionedListCacheMixin
from common.cache_utils import SCHOOL_LIST_CACHE_VERSION_KEY
from common.models import ImportJob
from common.pagination import BoundedPageNumberPagination
from common.permissions import IsParent, IsSchoolAdmin, RoleScopedPermission
from common.scoping import ScopedQuerysetMixin
from common.throttles import StudentClaimIPThrottle, StudentClaimUserThrottle

from .import_service import TEMPLATE_COLUMNS
from .models import City, School, Student
from .serializers import (
    CitySerializer,
    SchoolSerializer,
    StudentApproveSerializer,
    StudentAssignParentSerializer,
    StudentClaimSerializer,
    StudentImportJobSerializer,
    StudentImportUploadSerializer,
    StudentManualAddSerializer,
    StudentSerializer,
)
from .tasks import commit_student_import, validate_student_import


class StudentPagination(BoundedPageNumberPagination):
    """Requirement 1: student list is paginated at 50 rows per page."""

    page_size = 50
    max_page_size = 50


class CityViewSet(
    VersionedListCacheMixin,
    ScopedQuerysetMixin,
    viewsets.ModelViewSet,
):
    cache_version_key = SCHOOL_LIST_CACHE_VERSION_KEY
    cache_namespace = "schools:cities"
    queryset = City.objects.order_by("name")
    serializer_class = CitySerializer
    permission_classes = [RoleScopedPermission]
    read_roles = ("BOSS", "ADMIN", "SCHOOL_ADMIN", "TEACHER", "PARENT")
    write_roles = ("BOSS",)

    scope_city_field = "id"
    scope_school_field = None
    scope_parent_field = None

    def scope_queryset(self, qs):
        user = self.request.user
        if not user or not user.is_authenticated:
            return qs.none()
        if user.role == "BOSS" or getattr(user, "is_superuser", False):
            return qs
        if user.city_id:
            return qs.filter(id=user.city_id)
        return qs.none()


class SchoolViewSet(
    VersionedListCacheMixin,
    ScopedQuerysetMixin,
    viewsets.ModelViewSet,
):
    cache_version_key = SCHOOL_LIST_CACHE_VERSION_KEY
    cache_namespace = "schools:schools"
    queryset = School.objects.select_related("city").order_by("name")
    serializer_class = SchoolSerializer
    permission_classes = [RoleScopedPermission]
    read_roles = ("BOSS", "ADMIN", "SCHOOL_ADMIN", "TEACHER", "PARENT")
    write_roles = ("BOSS", "ADMIN")

    scope_city_field = "city_id"
    scope_school_field = "id"
    scope_parent_field = None

    def scope_queryset(self, qs):
        user = self.request.user
        if not user or not user.is_authenticated:
            return qs.none()
        if user.role == "PARENT":
            if user.school_id:
                return qs.filter(id=user.school_id)
            return qs.none()
        return super().scope_queryset(qs)


class StudentViewSet(ScopedQuerysetMixin, viewsets.ModelViewSet):
    """
    Student management (Prompt 4, requirements 1, 2, 4, 5, 6).

    Scoping (Prompt 3):
      - BOSS            : all students
      - ADMIN           : WHERE city_id = user.city_id
      - SCHOOL_ADMIN /
        TEACHER         : WHERE school_id = user.school_id
      - PARENT          : WHERE parent_id = user.id

    List endpoint: 50 rows/page, `?search=` on name + GR number (served by the
    functional `UPPER(...) varchar_pattern_ops` indexes), `?class_name=` /
    `?section=` filters served by the composite (school, class, section) index.
    """

    queryset = Student.objects.select_related("school", "parent").order_by(
        "school_id", "class_name", "section", "name"
    )
    serializer_class = StudentSerializer
    pagination_class = StudentPagination
    permission_classes = [RoleScopedPermission]
    read_roles = ("BOSS", "ADMIN", "SCHOOL_ADMIN", "TEACHER", "PARENT")
    # Requirement 2: Teachers and School Admins add students by form.
    write_roles = ("BOSS", "ADMIN", "SCHOOL_ADMIN", "TEACHER")

    scope_city_field = "city_id"
    scope_school_field = "school_id"
    scope_parent_field = "parent_id"

    # ------------------------------------------------------------------ #
    # Filtering / search
    # ------------------------------------------------------------------ #
    def get_queryset(self):
        qs = super().get_queryset()
        params = self.request.query_params

        school_id = params.get("school")
        class_name = params.get("class_name") or params.get("class")
        section = params.get("section")
        search = (params.get("search") or params.get("q") or "").strip()
        approval_status = params.get("approval_status")
        pending = params.get("pending")

        if school_id:
            qs = qs.filter(school_id=school_id)
        if class_name:
            qs = qs.filter(class_name=class_name)
        if section:
            qs = qs.filter(section=section)
        if search:
            # Prefix search -> `UPPER(name) LIKE UPPER('term%')`, served by the
            # idx_student_name_upper / idx_student_gr_upper expression indexes.
            # No leading wildcard, so no sequential scan and no per-row query.
            qs = qs.filter(
                Q(name__istartswith=search) | Q(gr_number__istartswith=search)
            )
        if approval_status:
            qs = qs.filter(approval_status=approval_status)
        if pending in ("1", "true", "True"):
            qs = qs.filter(approval_status=Student.ApprovalStatus.PENDING)
        return qs

    def list(self, request, *args, **kwargs):
        # A parent whose verified number appears on a roster sees those children
        # without claiming anything (one UPDATE, usually touching no rows).
        if request.user.is_authenticated and request.user.role == User.Role.PARENT:
            from .linking import link_children_by_phone

            link_children_by_phone(request.user)
        return super().list(request, *args, **kwargs)

    def check_permissions(self, request):
        if request.user.is_authenticated and request.user.role == User.Role.PARENT:
            if request.method not in ("GET", "HEAD", "OPTIONS") and self.action not in ("claim", "manual_add"):
                raise PermissionDenied("Student details are view-only for parents. Contact your school to correct this.")
        super().check_permissions(request)

    def create(self, request, *args, **kwargs):
        if request.user.role == User.Role.PARENT:
            raise PermissionDenied("Parents cannot create students directly. Contact your school to correct this.")
        return super().create(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        if request.user.role == User.Role.PARENT:
            raise PermissionDenied("Student details are view-only for parents. Contact your school to correct this.")
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        if request.user.role == User.Role.PARENT:
            raise PermissionDenied("Student details are view-only for parents. Contact your school to correct this.")
        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        if request.user.role == User.Role.PARENT:
            raise PermissionDenied("Parents cannot delete students. Contact your school to correct this.")
        return super().destroy(request, *args, **kwargs)

    def perform_create(self, serializer):
        if self.request.user.role == User.Role.PARENT:
            raise PermissionDenied("Parents cannot create students directly. Contact your school to correct this.")
        serializer.save()

    def perform_update(self, serializer):
        if self.request.user.role == User.Role.PARENT:
            raise PermissionDenied("Student details are view-only for parents. Contact your school to correct this.")
        serializer.save()

    def perform_destroy(self, instance):
        if self.request.user.role == User.Role.PARENT:
            raise PermissionDenied("Parents cannot delete students. Contact your school to correct this.")
        instance.delete()

    # ------------------------------------------------------------------ #
    # Permissions / throttles for the parent-facing actions
    # ------------------------------------------------------------------ #
    def get_permissions(self):
        if self.action in ("claim", "manual_add"):
            return [IsParent()]
        if self.action in ("approve", "assign_parent"):
            return [IsSchoolAdmin()]
        return super().get_permissions()

    def get_throttles(self):
        if self.action == "claim":
            # Requirement 6: GR numbers cannot be brute-forced.
            return [StudentClaimUserThrottle(), StudentClaimIPThrottle()]
        if self.action == "manual_add":
            return [StudentClaimUserThrottle(), StudentClaimIPThrottle()]
        return super().get_throttles()

    # ------------------------------------------------------------------ #
    # Requirement 4: parent claims a child
    # ------------------------------------------------------------------ #
    @action(detail=False, methods=["post"], url_path="claim")
    def claim(self, request):
        serializer = StudentClaimSerializer(
            data=request.data, context={"request": request}
        )
        serializer.is_valid(raise_exception=True)
        student = serializer.save()
        return Response(
            StudentSerializer(student, context={"request": request}).data,
            status=status.HTTP_200_OK,
        )

    # ------------------------------------------------------------------ #
    # Requirement 4: parent adds a child manually when there is no roster yet
    # ------------------------------------------------------------------ #
    @action(detail=False, methods=["post"], url_path="manual-add")
    def manual_add(self, request):
        serializer = StudentManualAddSerializer(
            data=request.data, context={"request": request}
        )
        serializer.is_valid(raise_exception=True)
        student = serializer.save()
        return Response(
            StudentSerializer(student, context={"request": request}).data,
            status=status.HTTP_201_CREATED,
        )

    # ------------------------------------------------------------------ #
    # Requirement 4: School Admin approves a pending (parent-added) student
    # ------------------------------------------------------------------ #
    @action(detail=True, methods=["post"], url_path="approve")
    def approve(self, request, pk=None):
        student = self.get_object()
        serializer = StudentApproveSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        for field in ("name", "gr_number", "class_name", "section", "gender"):
            if data.get(field):
                setattr(student, field, data[field])
        if "date_of_birth" in data:
            student.date_of_birth = data["date_of_birth"]

        if "parent" in request.data:
            parent_id = data.get("parent")
            if parent_id:
                student.parent = get_object_or_404(
                    User, pk=parent_id, role=User.Role.PARENT
                )
            else:
                student.parent = None
        student.approval_status = Student.ApprovalStatus.APPROVED
        student.save()
        return Response(
            StudentSerializer(student, context={"request": request}).data
        )

    # ------------------------------------------------------------------ #
    # Requirement 5: School Admin override of the parent link
    # ------------------------------------------------------------------ #
    @action(detail=True, methods=["post"], url_path="assign-parent")
    def assign_parent(self, request, pk=None):
        student = self.get_object()
        serializer = StudentAssignParentSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        parent_id = serializer.validated_data.get("parent")

        if parent_id:
            parent = get_object_or_404(User, pk=parent_id, role=User.Role.PARENT)
            student.parent = parent
        else:
            student.parent = None
        student.approval_status = Student.ApprovalStatus.APPROVED
        student.save(update_fields=["parent", "approval_status", "updated_at"])
        return Response(
            StudentSerializer(student, context={"request": request}).data
        )


class StudentImportViewSet(ScopedQuerysetMixin, viewsets.ReadOnlyModelViewSet):
    """
    Bulk student import (Prompt 4, requirement 3).

    POST   /api/student-imports/                upload file -> job id (202)
    GET    /api/student-imports/                list jobs (paginated)
    GET    /api/student-imports/{id}/           poll status + preview
    POST   /api/student-imports/{id}/confirm/   confirm -> second background job
                                                 body: {"mark_missing_as_left": true}
                                                 when the file is the school's
                                                 complete list (yearly upload)
    GET    /api/student-imports/{id}/report/    CSV of errors + duplicates
    GET    /api/student-imports/template/       downloadable template (.xlsx/.csv)
    """

    queryset = ImportJob.objects.select_related("school", "uploaded_by").order_by(
        "-created_at"
    )
    serializer_class = StudentImportJobSerializer
    permission_classes = [IsSchoolAdmin]
    parser_classes = [MultiPartParser, FormParser, JSONParser]
    pagination_class = BoundedPageNumberPagination

    scope_city_field = "city_id"
    scope_school_field = "school_id"
    scope_parent_field = None

    def get_queryset(self):
        qs = super().get_queryset().filter(job_type=ImportJob.JobType.STUDENT_IMPORT)
        school_id = self.request.query_params.get("school")
        if school_id:
            qs = qs.filter(school_id=school_id)
        return qs

    # ------------------------------------------------------------------ #
    # Upload (returns a job id immediately - nothing heavy in the request)
    # ------------------------------------------------------------------ #
    def create(self, request, *args, **kwargs):
        serializer = StudentImportUploadSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        upload = serializer.validated_data["file"]
        user = request.user

        school = serializer.validated_data.get("school") or None
        if user.role in ("SCHOOL_ADMIN", "TEACHER"):
            school = School.objects.select_related("city").get(pk=user.school_id)
        elif school is None:
            raise ValidationError({"school": "School is required."})
        elif user.role == "ADMIN" and school.city_id != user.city_id:
            raise PermissionDenied(
                "Admins can only import students for schools in their city."
            )

        job = ImportJob.objects.create(
            job_type=ImportJob.JobType.STUDENT_IMPORT,
            status=ImportJob.Status.PENDING,
            school=school,
            city_id=school.city_id,
            uploaded_by=user,
            original_filename=(upload.name or "upload")[:255],
            file=upload,
        )

        # Rule P5: validation runs in a Celery worker, never in this request.
        validate_student_import.delay(str(job.id))

        job.refresh_from_db()
        payload = StudentImportJobSerializer(job, context={"request": request}).data
        return Response(payload, status=status.HTTP_202_ACCEPTED)

    # ------------------------------------------------------------------ #
    # Confirm -> batched insert job
    # ------------------------------------------------------------------ #
    @action(detail=True, methods=["post"])
    def confirm(self, request, pk=None):
        job = self.get_object()
        if job.status == ImportJob.Status.FAILED:
            raise ValidationError(
                {"detail": f"This job failed and cannot be imported: {job.error_message}"}
            )
        if job.status in (
            ImportJob.Status.IMPORTING,
            ImportJob.Status.COMPLETED,
        ):
            raise ValidationError(
                {"detail": f"This job is already {job.status.lower()}."}
            )
        if job.status != ImportJob.Status.PREVIEW_READY:
            raise ValidationError(
                {"detail": "The file is still being validated. Try again shortly."}
            )

        raw_flag = request.data.get("mark_missing_as_left", False)
        mark_missing = str(raw_flag).lower() in ("1", "true", "yes", "on")
        job.preview = {**(job.preview or {}), "mark_missing_as_left": mark_missing}
        job.save(update_fields=["preview", "updated_at"])

        commit_student_import.delay(str(job.id))
        job.refresh_from_db()
        return Response(
            StudentImportJobSerializer(job, context={"request": request}).data,
            status=status.HTTP_202_ACCEPTED,
        )

    # ------------------------------------------------------------------ #
    # CSV report of rejected rows (bounded payload for the UI, Rule P2)
    # ------------------------------------------------------------------ #
    @action(detail=True, methods=["get"], url_path="report")
    def report(self, request, pk=None):
        job = self.get_object()
        preview = job.preview or {}

        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = (
            f'attachment; filename="import-{job.id}-report.csv"'
        )
        writer = csv.writer(response)
        writer.writerow(["row", "name", "gr_number", "class", "section", "reason"])
        for row in preview.get("duplicates", []):
            writer.writerow(
                [
                    row.get("row", ""),
                    row.get("name", ""),
                    row.get("gr_number", ""),
                    row.get("class", ""),
                    row.get("section", ""),
                    row.get("reason", "Duplicate"),
                ]
            )
        for row in preview.get("errors", []):
            reasons = "; ".join(f"{k}: {v}" for k, v in (row.get("errors") or {}).items())
            writer.writerow(
                [
                    row.get("row", ""),
                    row.get("name", ""),
                    row.get("gr_number", ""),
                    "",
                    "",
                    reasons,
                ]
            )
        return response

    # ------------------------------------------------------------------ #
    # Downloadable template with the exact columns
    # ------------------------------------------------------------------ #
    @action(detail=False, methods=["get"], url_path="template")
    def template(self, request):
        fmt = (
            request.query_params.get("file_format")
            or request.query_params.get("format")
            or "xlsx"
        ).lower()
        if fmt not in ("xlsx", "csv"):
            raise ValidationError({"format": "Use ?format=xlsx or ?format=csv."})

        if fmt == "csv":
            response = HttpResponse(content_type="text/csv")
            response["Content-Disposition"] = (
                'attachment; filename="student-import-template.csv"'
            )
            writer = csv.writer(response)
            writer.writerow(TEMPLATE_COLUMNS)
            return response

        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Students"
        sheet.append(TEMPLATE_COLUMNS)
        for index, width in enumerate([24, 16, 10, 10, 12, 16], start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"

        instructions = workbook.create_sheet("Instructions")
        instructions.append(["Column", "Required", "Accepted values"])
        instructions.append(["name", "Yes", "Student full name (max 150 chars)"])
        instructions.append(["gr_number", "Yes", "Unique per school (max 50 chars)"])
        instructions.append(["class", "Yes", "Class/standard, e.g. 5, 10-A (max 30 chars)"])
        instructions.append(["section", "Yes", "Section, e.g. A, B (max 20 chars)"])
        instructions.append(["gender", "Yes", "M / F / OTHER (Male, Female, Boy, Girl accepted)"])
        instructions.append(["date_of_birth", "No", "YYYY-MM-DD (used when a parent claims the child)"])
        instructions.append([])
        instructions.append(["Limits", "", "Max 5 MB file, max 5,000 rows per upload"])
        instructions.append(["Duplicates", "", "Rows whose GR number already exists are skipped"])
        for index, width in enumerate([16, 12, 60], start=1):
            instructions.column_dimensions[get_column_letter(index)].width = width

        buffer = io.BytesIO()
        workbook.save(buffer)
        response = HttpResponse(
            buffer.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
        )
        response["Content-Disposition"] = (
            'attachment; filename="student-import-template.xlsx"'
        )
        return response
